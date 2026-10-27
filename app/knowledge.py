
from __future__ import annotations
import argparse
import csv
import hashlib
import json
import re
import shutil
import sqlite3
import subprocess
import tempfile
import unicodedata
from contextlib import closing
from difflib import SequenceMatcher
from pathlib import Path
from pypdf import PdfReader

VERSION = 'kb-1'
DEFAULT_DB = Path(__file__).resolve().parents[1] / 'work' / 'knowledge.sqlite3'

def norm(text):
    return re.sub(r'\s+', ' ', unicodedata.normalize('NFKC', text).casefold()).strip()

def digest(text):
    return hashlib.sha256(text.encode('utf-8')).hexdigest()

def descriptor_grid(text):
    start=re.search(r'\b(?:band|level)\s+8\b',text)
    if not start or not re.search(r'\b(?:band|level)\s+[1-7]\b',text): return None
    grid=text[start.start():]
    grid=re.split(r'papacambridge|downloaded from|licensed for hosting',grid)[0]
    
    grid=re.sub(r'\bband\b','level',grid)
    grid=re.sub(r'\s*/\s*','/',grid)
    return norm(grid.replace('\ufffd',' ').replace('•',' '))

def db_open(path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    c = sqlite3.connect(path, timeout=45)
    c.row_factory = sqlite3.Row
    c.executescript('''
    PRAGMA journal_mode=WAL;
    CREATE TABLE IF NOT EXISTS documents(id TEXT PRIMARY KEY, manifest TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS pages(resource_id TEXT, page INTEGER, raw TEXT, normalized TEXT,
      method TEXT, flags TEXT, PRIMARY KEY(resource_id,page));
    CREATE TABLE IF NOT EXISTS passages(id TEXT PRIMARY KEY, resource_id TEXT, page INTEGER,
      start INTEGER, end INTEGER, text TEXT, provenance TEXT);
    CREATE VIRTUAL TABLE IF NOT EXISTS search USING fts5(passage_id UNINDEXED,text);
    CREATE TABLE IF NOT EXISTS notes(id TEXT PRIMARY KEY, resource_id TEXT, page INTEGER, record TEXT);
    CREATE TABLE IF NOT EXISTS rubrics(id TEXT PRIMARY KEY, record TEXT);
    CREATE TABLE IF NOT EXISTS calibration(id TEXT PRIMARY KEY, record TEXT);
    ''')
    return c

def quality(text):
    flags = []
    if not text.strip(): flags.append('empty_page')
    elif len(re.findall(r'\w', text)) < 35: flags.append('sparse_text')
    if '\ufffd' in text or '\x00' in text: flags.append('broken_encoding')
    lines = [norm(x) for x in text.splitlines() if len(x.strip()) > 30]
    if lines and len(set(lines)) < len(lines)*.7: flags.append('repeated_lines')
    if len(re.findall(r'\b\w\b', text)) > max(20, len(text.split())*.4):
        flags.append('possible_reading_order_or_encoding_issue')
    return flags

def ocr_image(path):
    executable = shutil.which('tesseract')
    if not executable: return '', 'ocr_unavailable'
    run = subprocess.run([executable, str(path), 'stdout'], capture_output=True, timeout=120)
    if run.returncode: return '', 'ocr_failed'
    return run.stdout.decode('utf-8', errors='replace'), 'tesseract'

def extract(path):
    
    suffix = path.suffix.lower()
    result = []
    if suffix == '.pdf':
        reader = PdfReader(path)
        for index, page in enumerate(reader.pages, 1):
            try:
                raw = page.extract_text() or ''
                flags = quality(raw)
            except Exception as exc:
                raw=''; flags=['native_page_extraction_failed',type(exc).__name__]
            text, method = raw, 'pypdf_native'
            if not raw.strip() or 'broken_encoding' in flags:
                renderer = shutil.which('pdftoppm')
                if renderer and shutil.which('tesseract'):
                    with tempfile.TemporaryDirectory() as temp:
                        target = Path(temp)/'page'
                        try:
                            subprocess.run([renderer, '-f', str(index), '-l', str(index), '-singlefile',
                                            '-r', '200', '-png', str(path), str(target)],
                                           check=True, capture_output=True, timeout=120)
                            replacement, ocr_method = ocr_image(target.with_suffix('.png'))
                            if replacement.strip(): text,method=replacement,ocr_method
                            else: flags.append(ocr_method if ocr_method.startswith('ocr_') else 'ocr_empty')
                        except (OSError,subprocess.SubprocessError) as exc:
                            flags.extend(['ocr_failed',type(exc).__name__])
                else: flags.append('ocr_unavailable')
            result.append({'page':index, 'raw':raw, 'text':text, 'method':method,
                           'flags':list(dict.fromkeys(flags + quality(text)))})
    elif suffix in {'.txt', '.md', '.csv'}:
        raw = path.read_text(encoding='utf-8-sig', errors='replace')
        result.append({'page':1,'raw':raw,'text':raw,'method':'text_utf8','flags':quality(raw)})
    elif suffix == '.docx':
        from docx import Document
        document = Document(path)
        raw = '\n'.join([p.text for p in document.paragraphs] +
                        ['\t'.join(cell.text for cell in row.cells) for table in document.tables for row in table.rows])
        result.append({'page':1,'raw':raw,'text':raw,'method':'python_docx','flags':['logical_page_not_print_page']})
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            for index, sheet in enumerate(workbook, 1):
                raw = sheet.title+'\n'+'\n'.join('\t'.join('' if x is None else str(x) for x in row) for row in sheet.values)
                result.append({'page':index,'raw':raw,'text':raw,'method':'openpyxl','flags':['sheet_not_print_page']})
        finally: workbook.close()
    elif suffix in {'.png','.jpg','.jpeg','.tif','.tiff','.bmp'}:
        raw, method = ocr_image(path)
        result.append({'page':1,'raw':raw,'text':raw,'method':method,'flags':quality(raw)+([method] if method.startswith('ocr_') else [])})
    else:
        result.append({'page':1,'raw':'','text':'','method':'unsupported','flags':['unsupported_format']})
    
    boundaries = {}
    for p in result:
        lines = p['text'].splitlines()
        for line in set(lines[:2]+lines[-2:]):
            if 0 < len(line.strip()) < 120: boundaries[line] = boundaries.get(line,0)+1
    repeated = {x for x,n in boundaries.items() if len(result)>2 and n>=max(3,len(result)*.5)}
    for p in result:
        p['normalized'] = norm('\n'.join(x for x in p['text'].splitlines() if x not in repeated))
        if any(x in p['text'].splitlines() for x in repeated): p['flags'].append('repeated_boundary_lines_excluded_from_search')
    return result

def metadata(path, text):
    
    head = text[:14000]
    low = norm(head)
    rules = [('examiner_report',r'^.{0,200}principal examiner reports?|^.{0,100}key messages'),
             ('syllabus',r'^.{0,300}syllabus overview|syllabus for examination'),
             ('cambridge_teaching_pack',r'^.{0,300}learner guide|using candidate responses in the classroom|teaching pack|lesson plan'),
             ('candidate_response',r'^.{0,500}(?:example|exemplar) candidate responses'),
             ('generic_mark_scheme',r'^.{0,700}mark scheme'),
             ('question_paper',r'^.{0,700}question paper|answer (?:one|two|all).*questions|instructions.*candidates')]
    category = next((kind for kind,pattern in rules if re.search(pattern,low)), 'unknown')
    evidence = 'content_rule' if category!='unknown' else 'unresolved'
    works = {'things fall apart':'Chinua Achebe','blues for an alabama sky':'Pearl Cleage',
             'children of wealth':None,'the road':None,'the bus':None}
    found = [w for w in works if re.search(r'\b'+re.escape(w)+r'\b',low)]
    work = found[0] if len(found)==1 else None
    if category=='unknown' and work:
        category='literary_notes'
        evidence='content_candidate_requires_review'
    name_match = re.search(r'0475[_ -]([msw])(\d{2})[_ -](?:ms|qp|er)[_ -]?(\d{2})?',path.name.lower())
    content_years = sorted({int(y) for y in re.findall(r'\b20(?:2[0-9]|30)\b',head[:5000])})
    year = 2000+int(name_match[2]) if name_match else (content_years[0] if len(content_years)==1 else None)
    component = name_match[3] if name_match else None
    content_component = re.search(r'0475\s*/\s*([1-4]\d)',head)
    conflict = bool(content_component and component and content_component[1]!=component)
    if content_component: component=content_component[1]
    if name_match and content_years and year not in content_years: conflict=True
    policy = category in {'generic_mark_scheme','question_specific_mark_scheme','assessment_criteria','syllabus'}
    active = policy and year in {2026,2027} and not conflict and not any(y>=2028 for y in content_years)
    
    
    known_primary = {
        '_oceanofpdf.com_blues_for_an_alabama_sky_-_pearl_cleage.pdf':('blues for an alabama sky','Pearl Cleage'),
        'things-fall-apart.pdf':('things fall apart','Chinua Achebe'),
        'children of wealth.md':('children of wealth',None),
        'the bus by arun.k.pdf':('the bus',None),
        'the roadpoem.pdf':('the road',None),
    }
    binding=known_primary.get(path.name.lower())
    if binding and binding[0] in norm(text):
        category='primary_literary_text'; work=binding[0]
        evidence='user_inventory_binding_with_content_title_match'
    return dict(category=category,classification_evidence=evidence,confidence='review_required',
                syllabus='0475' if '0475' in head else None,year=year,content_years=content_years,
                component=component,session=name_match[1] if name_match else None,work=work,
                author=binding[1] if binding and category=='primary_literary_text' else works.get(work),question_number=None,metadata_conflict=conflict,
                active_policy=bool(active and '0475' in head),verification_status='unverified_extraction')

def raw_annotations(path, pages):
    from .core import annotation_records
    reconstructed = annotation_records(path)
    reader=PdfReader(path)
    by_page={}
    for r in reconstructed: by_page.setdefault(r['page'],[]).append(r)
    for page_no,page in enumerate(reader.pages,1):
        for ref in page.get('/Annots',[]):
            obj=ref.get_object()
            kind=str(obj.get('/Subtype','/Unknown')).lstrip('/')
            if kind=='Link': continue
            rect=[float(x) for x in obj.get('/Rect',[])]
            match=next((r for r in by_page.get(page_no,[]) if r['subtype']==kind and
                        (r.get('comment') or '')==str(obj.get('/Contents','')).strip()),{})
            if match: by_page[page_no].remove(match)
            phrase=match.get('highlighted_text') or ''
            page_text=next(p['normalized'] for p in pages if p['page']==page_no)
            position=page_text.find(norm(phrase)) if phrase else -1
            def ref_id(value):
                return str(getattr(value,'idnum',value)) if value is not None else None
            comment=str(obj.get('/Contents',''))
            popup=obj.get('/Popup')
            popup_text=str(popup.get_object().get('/Contents','')) if popup else (comment if kind=='Popup' else '')
            labels=[label for label in ['language','imagery','symbolism','theme','character','relationship',
                    'structure','tone','context','motif'] if label in comment.lower()]
            yield {**match,'page':page_no,'type':kind,'object_id':ref_id(ref),'rect':rect,
                   'quadpoints':[float(x) for x in obj.get('/QuadPoints',[])],
                   'original_comment':comment,'popup_text':popup_text,'author':str(obj.get('/T','')),
                   'created':str(obj.get('/CreationDate','')),'modified':str(obj.get('/M','')),
                   'popup_ref':ref_id(obj.raw_get('/Popup') if '/Popup' in obj else None),
                   'parent_ref':ref_id(obj.raw_get('/Parent') if '/Parent' in obj else None),
                   'reply_ref':ref_id(obj.raw_get('/IRT') if '/IRT' in obj else None),
                   'colour':[float(x) for x in obj.get('/C',[])],
                   'primary_span':{'start':position,'end':position+len(norm(phrase))} if position>=0 else None,
                   'unresolved':position<0,'classification':{'method':'keyword_rules','labels':labels or ['other']},
                   'provenance':'human_annotation','chapter_act_scene':None}

def build(root, db_path=DEFAULT_DB, refresh_annotations=False):
    root=Path(root).resolve()
    if not root.is_dir(): raise ValueError('Resource folder does not exist')
    manifests=[]
    canonical_hashes={}; text_hashes={}; signatures=[]
    with closing(db_open(db_path)) as c:
        for path in sorted(p for p in root.rglob('*') if p.is_file() and p.name.lower()!='desktop.ini'):
            rid=digest(str(path.resolve()))
            checksum=hashlib.sha256(path.read_bytes()).hexdigest()
            old=c.execute('SELECT manifest FROM documents WHERE id=?',(rid,)).fetchone()
            cached=json.loads(old[0]) if old else None
            if cached and cached['checksum']==checksum and cached.get('version')==VERSION and not cached.get('errors'):
                m=cached
                if refresh_annotations and m.get('annotation_count'):
                    pages=[dict(r) for r in c.execute('SELECT page,normalized FROM pages WHERE resource_id=?',(rid,))]
                    records=list(raw_annotations(path,pages))
                    c.execute('DELETE FROM notes WHERE resource_id=?',(rid,))
                    for n,note in enumerate(records):
                        note['work']=m['work']; span=note.get('primary_span')
                        note['passage_id']=f'{rid}:{note["page"]}:{(span["start"]//850)*850}' if span else None
                        c.execute('INSERT INTO notes VALUES(?,?,?,?)',(f'{rid}:{n}',rid,note['page'],json.dumps(note,ensure_ascii=False)))
                    m['annotation_count']=len(records)
                    print(json.dumps({'annotations_refreshed':path.name,'count':len(records)}),flush=True)
                    c.commit()
                canonical_hashes.setdefault(checksum,m['canonical_resource_id'])
                if m['text_hash']: text_hashes.setdefault(m['text_hash'],m['canonical_resource_id'])
                manifests.append(m)
                signatures.append((rid,m.get('signature',[]),m['category']))
                continue
            errors=[]
            try: pages=extract(path)
            except Exception as exc:
                pages=[dict(page=1,raw='',normalized='',method='failed',flags=['extraction_failed'])]
                errors.append(type(exc).__name__+': '+str(exc))
            text='\n'.join(p['normalized'] for p in pages)
            m=metadata(path,'\n'.join(p['raw'] for p in pages))
            text_hash=digest(text) if text.strip() else None
            canonical=canonical_hashes.get(checksum) or text_hashes.get(text_hash) or rid
            duplicate='exact_file' if checksum in canonical_hashes else 'normalized_text' if text_hash in text_hashes else 'unique'
            signature=sorted(set(re.findall(r'\b\w{4,}\b',text)))
            near=[]
            for other,other_sig,category in signatures:
                if duplicate!='unique' or category!=m['category'] or len(signature)<30: continue
                a,b=set(signature),set(other_sig)
                similarity=len(a & b)/max(1,len(a | b))
                if similarity>=.9: near.append({'resource_id':other,'similarity':round(similarity,4),'status':'review_required_not_merged'})
            m.update(id=rid,path=str(path),filename=path.name,file_type=path.suffix.lower(),checksum=checksum,
                     page_count=len(pages),text_hash=text_hash,canonical_resource_id=canonical,duplicate_status=duplicate,
                     fuzzy_duplicate_candidates=near,signature=signature,version=VERSION,errors=errors,
                     extraction_method=sorted({p['method'] for p in pages}),
                     extraction_quality=[{'page':p['page'],'flags':p['flags']} for p in pages],annotation_count=0)
            c.execute('DELETE FROM search WHERE passage_id IN (SELECT id FROM passages WHERE resource_id=?)',(rid,))
            for table in ('pages','passages','notes'): c.execute(f'DELETE FROM {table} WHERE resource_id=?',(rid,))
            for p in pages:
                c.execute('INSERT INTO pages VALUES(?,?,?,?,?,?)',(rid,p['page'],p['raw'],p['normalized'],p['method'],json.dumps(p['flags'])))
                for start in range(0,len(p['normalized']),850):
                    end=min(start+1100,len(p['normalized']))
                    pid=f'{rid}:{p["page"]}:{start}'
                    c.execute('INSERT INTO passages VALUES(?,?,?,?,?,?,?)',(pid,rid,p['page'],start,end,p['normalized'][start:end],m['category']))
                    if canonical==rid: c.execute('INSERT INTO search VALUES(?,?)',(pid,p['normalized'][start:end]))
            if path.suffix.lower()=='.pdf':
                try:
                    
                    for n,note in enumerate(raw_annotations(path,pages)):
                        note['work']=m['work']
                        span=note.get('primary_span')
                        note['passage_id']=f'{rid}:{note["page"]}:{(span["start"]//850)*850}' if span else None
                        c.execute('INSERT INTO notes VALUES(?,?,?,?)',(f'{rid}:{n}',rid,note['page'],json.dumps(note,ensure_ascii=False)))
                        m['annotation_count']+=1
                except Exception as exc: m['errors'].append('annotation_extraction: '+str(exc))
            m['annotation_presence']=m['annotation_count']>0
            c.execute('INSERT OR REPLACE INTO documents VALUES(?,?)',(rid,json.dumps(m,ensure_ascii=False)))
            canonical_hashes.setdefault(checksum,canonical)
            if text_hash: text_hashes.setdefault(text_hash,canonical)
            signatures.append((rid,signature,m['category']))
            manifests.append(m)
            c.commit()
            print(json.dumps({'file':path.name,'pages':len(pages),'annotations':m['annotation_count'],'errors':m['errors']}),flush=True)
        
        
        
        canonical_hashes={}; text_hashes={}
        c.execute('DELETE FROM search')
        for m in manifests:
            raw='\n'.join(r[0] for r in c.execute('SELECT raw FROM pages WHERE resource_id=? ORDER BY page',(m['id'],)))
            m.update(metadata(Path(m['path']),raw))
            records=[(r['id'],json.loads(r['record'])) for r in c.execute('SELECT id,record FROM notes WHERE resource_id=?',(m['id'],))]
            objects={note['object_id']:note for _,note in records}
            for note_id,note in records:
                note['work']=m['work']
                popup=objects.get(note.get('popup_ref'))
                note['popup_text']=popup['original_comment'] if popup else (note['original_comment'] if note['type']=='Popup' else note.get('popup_text',''))
                c.execute('UPDATE notes SET record=? WHERE id=?',(json.dumps(note,ensure_ascii=False),note_id))
            canonical=canonical_hashes.get(m['checksum']) or text_hashes.get(m['text_hash']) or m['id']
            m['duplicate_status']='exact_file' if m['checksum'] in canonical_hashes else 'normalized_text' if m['text_hash'] in text_hashes else 'unique'
            m['canonical_resource_id']=canonical
            canonical_hashes.setdefault(m['checksum'],canonical)
            if m['text_hash']: text_hashes.setdefault(m['text_hash'],canonical)
            c.execute('UPDATE documents SET manifest=? WHERE id=?',(json.dumps(m,ensure_ascii=False),m['id']))
            if canonical==m['id']:
                c.execute('INSERT INTO search SELECT id,text FROM passages WHERE resource_id=?',(m['id'],))
        active={m['id'] for m in manifests}
        for row in c.execute('SELECT id FROM documents').fetchall():
            if row[0] not in active:
                c.execute('DELETE FROM search WHERE passage_id IN (SELECT id FROM passages WHERE resource_id=?)',(row[0],))
                for table in ('pages','passages','notes'): c.execute(f'DELETE FROM {table} WHERE resource_id=?',(row[0],))
                c.execute('DELETE FROM documents WHERE id=?',(row[0],))
        c.execute('DELETE FROM rubrics'); c.execute('DELETE FROM calibration')
        for m in manifests:
            if m['active_policy']:
                for page in c.execute('SELECT * FROM pages WHERE resource_id=?',(m['id'],)).fetchall():
                    grid=descriptor_grid(page['normalized'])
                    if grid:
                        
                        key=digest(grid)
                        old=c.execute('SELECT record FROM rubrics WHERE id=?',(key,)).fetchone()
                        record=json.loads(old[0]) if old else {'descriptor_text':grid,'sources':[],'status':'extracted_descriptor_requires_review'}
                        record['sources'].append({'resource_id':m['id'],'page':page['page'],'year':m['year'],'component':m['component']})
                        c.execute('INSERT OR REPLACE INTO rubrics VALUES(?,?)',(key,json.dumps(record)))
            if m['category'] in {'candidate_response','validation_response'}:
                c.execute('INSERT INTO calibration VALUES(?,?)',(m['id'],json.dumps({'resource_id':m['id'],
                    'canonical_resource_id':m['canonical_resource_id'],'status':'needs_question_response_mark_alignment',
                    'teacher_mark':None,'eligible_for_training':False,'split_group':m['canonical_resource_id']})))
        c.commit()
        return {'resources':len(manifests),'annotations':c.execute('SELECT count(*) FROM notes').fetchone()[0],
                'pages':c.execute('SELECT count(*) FROM pages').fetchone()[0],
                'rubrics':c.execute('SELECT count(*) FROM rubrics').fetchone()[0],
                'issues':[{'file':m['filename'],'errors':m['errors']} for m in manifests if m['errors']],
                'ocr_unavailable_pages':sum('ocr_unavailable' in p['flags'] for m in manifests for p in m['extraction_quality'])}

def retrieve(query, work=None, component=None, db_path=DEFAULT_DB, limit=8):
    if not Path(db_path).exists(): return []
    terms=list(dict.fromkeys(re.findall(r'\b\w{4,}\b',norm(query))))[:40]
    if not terms: return []
    with closing(db_open(db_path)) as c:
        rows=c.execute('''SELECT p.*,d.manifest FROM search JOIN passages p ON p.id=search.passage_id
          JOIN documents d ON d.id=p.resource_id WHERE search MATCH ? ORDER BY bm25(search) LIMIT 150''',
          (' OR '.join('"'+t+'"' for t in terms),)).fetchall()
        output=[]
        for row in rows:
            m=json.loads(row['manifest'])
            if m['metadata_conflict']: continue
            if m['category']=='unknown': continue
            if m['category'] in {'generic_mark_scheme','syllabus','assessment_criteria','question_specific_mark_scheme'}:
                if not m['active_policy'] or (component and m['component'] and str(component)!=m['component']): continue
            elif (m['year'] and m['year']>=2028) or any(y>=2028 for y in m['content_years']): continue
            if work and m['category'] in {'primary_literary_text','literary_notes','annotated_literary_text'}:
                if not m['work'] or norm(work)!=norm(m['work']): continue
            output.append({'passage_id':row['id'],'resource_id':row['resource_id'],'filename':m['filename'],
                           'page':row['page'],'text':row['text'],'provenance':m['category'],
                           'verification_status':m['verification_status']})
            if len(output)>=limit: break
        return output

def export(db_path, target):
    target=Path(target); target.mkdir(parents=True,exist_ok=True)
    with closing(db_open(db_path)) as c:
        for table in ('documents','pages','passages','notes','rubrics','calibration'):
            with (target/(table+'.jsonl')).open('w',encoding='utf-8') as f:
                for row in c.execute('SELECT * FROM '+table):
                    data=dict(row)
                    for key in ('record','manifest'):
                        if key in data: data[key]=json.loads(data[key])
                    f.write(json.dumps(data,ensure_ascii=False)+'\n')

if __name__=='__main__':
    from .core import SOURCE_ROOT
    parser=argparse.ArgumentParser()
    parser.add_argument('--root',type=Path,default=SOURCE_ROOT)
    parser.add_argument('--db',type=Path,default=DEFAULT_DB)
    parser.add_argument('--export',type=Path)
    parser.add_argument('--refresh-annotations',action='store_true')
    args=parser.parse_args()
    print(json.dumps(build(args.root,args.db,args.refresh_annotations)),flush=True)
    if args.export: export(args.db,args.export)
